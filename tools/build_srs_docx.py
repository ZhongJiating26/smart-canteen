from __future__ import annotations

import re
import sys
import zipfile
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "docs" / "A智能食堂系统需求规格说明书.md"

BLACK = RGBColor(0, 0, 0)
LIGHT_FILL = "F2F4F7"
BORDER = "808080"


def set_run_font(run, name="SimSun", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color="D9E2F3", size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[min(idx, len(widths) - 1)]))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5

    for style_name, size, before, after in [
        ("Heading 1", 16, 12, 6),
        ("Heading 2", 14, 10, 4),
        ("Heading 3", 12, 8, 2),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("软件需求规格说明书")
    set_run_font(r, name="SimHei", size=12, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(10)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(title_text)
    set_run_font(r, name="SimHei", size=22, bold=True)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(18)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Smart Canteen System Software Requirements Specification")
    set_run_font(r, name="Times New Roman", size=12)

    meta = doc.add_table(rows=5, cols=2)
    meta.autofit = False
    rows = [
        ("文档版本", "V1.0"),
        ("文档状态", "正式版"),
        ("编制日期", date.today().isoformat()),
        ("适用范围", "校园/园区智能食堂系统"),
        ("文档用途", "系统架构设计、详细设计、开发实现与测试验收依据"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
    format_table(meta, [1800, 7560], header=False)

    doc.add_paragraph()
    toc_title = doc.add_paragraph()
    toc_title.style = "Heading 1"
    toc_title.add_run("目录")
    toc = doc.add_paragraph()
    toc.paragraph_format.space_after = Pt(4)
    r = toc.add_run("请在 Word 中右键此处选择“更新域”，生成或刷新自动目录。")
    set_run_font(r, size=10.5, italic=True)
    add_toc_field(toc)
    doc.add_section(WD_SECTION.NEW_PAGE)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在更新域后显示"
    text_run = OxmlElement("w:r")
    text_run.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    paragraph._p.append(text_run)
    run._r.append(end)


def format_table(table, widths, header=True):
    table.autofit = False
    set_table_borders(table)
    set_table_width(table, widths)
    for ridx, row in enumerate(table.rows):
        if ridx == 0 and header:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tr_pr.append(tbl_header)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ridx == 0 and header:
                set_cell_shading(cell, LIGHT_FILL)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.10
                for run in p.runs:
                    set_run_font(run, size=9.5, bold=(ridx == 0 and header))


def parse_inline(paragraph, text):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\"[^\"]+\")", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=9.5)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def is_table_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells and all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells)


def split_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def table_widths(col_count):
    if col_count == 2:
        return [2160, 7200]
    if col_count == 3:
        return [1800, 3060, 4500]
    if col_count == 4:
        return [1560, 2340, 2340, 3120]
    if col_count == 5:
        return [1200, 2040, 2040, 2040, 2040]
    width = 9360 // col_count
    return [width] * col_count


def add_markdown_table(doc, rows):
    clean_rows = [r for idx, r in enumerate(rows) if idx != 1 or not is_table_separator(r)]
    data = [split_table_row(r) for r in clean_rows]
    col_count = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=col_count)
    for r_idx, row in enumerate(data):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx] if c_idx < len(row) else ""
    format_table(table, table_widths(col_count), header=True)


def build(source=DEFAULT_SOURCE):
    source = Path(source)
    if not source.is_absolute():
        source = ROOT / source
    output = source.with_suffix(".docx")
    title_text = source.stem

    doc = Document()
    setup_styles(doc)
    add_cover(doc, title_text)

    header = doc.sections[-1].header.paragraphs[0]
    header.text = title_text
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, size=9)
    footer = doc.sections[-1].footer.paragraphs[0]
    add_page_number(footer)

    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or set(stripped) <= {"-"}:
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                i += 1
                continue
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            p.add_run(text)
            for run in p.runs:
                set_run_font(run, bold=True)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            parse_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.font.italic = True
                run.font.color.rgb = BLACK
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            parse_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            parse_inline(p, stripped[2:].strip())
            i += 1
            continue

        p = doc.add_paragraph()
        parse_inline(p, stripped)
        i += 1

    scrub_docx_metadata(doc)
    doc.save(output)
    patch_update_fields(output)
    return output


def scrub_docx_metadata(doc):
    props = doc.core_properties
    props.author = "Smart Canteen Project Team"
    props.last_modified_by = "Smart Canteen Project Team"
    props.title = "A智能食堂系统需求规格说明书"
    props.subject = "软件需求规格说明书"
    props.keywords = "智能食堂, 需求规格说明书, SRS"
    props.comments = ""


def patch_update_fields(path):
    settings_name = "word/settings.xml"
    temp_path = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == settings_name:
                xml = data.decode("utf-8")
                if "w:updateFields" not in xml:
                    xml = xml.replace(
                        "</w:settings>",
                        '<w:updateFields w:val="true"/></w:settings>',
                    )
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    temp_path.replace(path)


if __name__ == "__main__":
    sources = sys.argv[1:] or [str(DEFAULT_SOURCE)]
    for source in sources:
        print(build(source))
